import docx
from schoolReport import getNamesGrades, coordinates, testDictionary

def testWriteSchoolReport(inputDocx):
    xlsxFile = "/home/alex/schoolReport/Notenliste-11-10-2021_04-02-21.xlsx"
    #names, allGrades = getNamesGrades.getNamesGrades(xlsxFile)
    names, allGrades = testDictionary.testDictionary(xlsxFile)

    coordinatesDocX = coordinates.coordinates()

    def getCoordinates(schoolSubj):
        """ get Coordinates of "note" for print to docX """
        #get Coordinates
        for i in range(len(coordinatesDocX)):

            if schoolSubj == coordinatesDocX[i][0]:
                c = coordinatesDocX[i][3]
                r = coordinatesDocX[i][2]

                print("coordinates of:  "+schoolSubj+" column: "+str(c)+", row: " +str(r))
                tableDocx = doc.tables[coordinatesDocX[i][4]]
                targetCell = tableDocx.cell(r, c)

                return targetCell

    #get name
    pupil = 0
    for name in names:

        #open .docx
        print(inputDocx)
        outputDocx = "output/"+name+"-saved-"+inputDocx
        savepath= outputDocx
        doc = docx.Document(inputDocx)

        #get note from allGrades
        for schoolSubject in allGrades:
            note = allGrades[schoolSubject]["Gesamtnote"][pupil]
            targetCell = getCoordinates(schoolSubject)
            if targetCell != None:
                #write to target
                targetCell.text = str(note)

        #get note from allGrades
        note = allGrades["7a - Deu"]["Ø Mitarbeit"][pupil]
        targetCell = getCoordinates("Deu_m")
        if targetCell != None:
            #write to target
            targetCell.text = str(note)

        #get note from allGrades
        note = allGrades["7a - Deu"]["Ø Klausur"][pupil]
        targetCell = getCoordinates("Deu_s")
        if targetCell != None:
            #write to target
            targetCell.text = str(note)


        pupil += 1



        #save to target .docX
        doc.save(savepath)
        print("Was saved in", savepath)


if __name__ == "__main__":

    docxFile = "Z 420 - Notenzeugnis für Schülerinnen und Schüler mit dem Förderbedarf Lernen (01.21)-1.docx"
    testWriteSchoolReport(docxFile)
