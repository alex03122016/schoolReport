import docx
from schoolReport import getNamesGrades, coordinates, testDictionary, writePupilDetailsToSchoolReport, testwritePupilDetailsToSchoolReport, writeParticipationToSchoolReport

def writeSchoolReport(inputDocx):
    """ watch out: the exported file from lehrmeister has to be saved once in libreoffice
    as xlsx file
    else it will throw the following error:
        TypeError: expected <class 'openpyxl.styles.fills.Fill'>
    """
    xlsxFile = "/home/alex/schoolReport/Notenliste-18-10-2021_08-34-50.xlsx"


    #names, allGrades = getNamesGrades.getNamesGrades(xlsxFile)
    names, allGrades = testDictionary.testDictionary(xlsxFile)

    coordinatesDocX = coordinates.coordinates()

    def getCoordinates(schoolSubj):
        """ get Coordinates of "note" for print to docX """
        #get Coordinates
        for i in range(len(coordinatesDocX)):

            if schoolSubj == coordinatesDocX[i][0]:
                c = coordinatesDocX[i][3]
                r = coordinatesDocX[i][2]

                print("coordinates of:  "+schoolSubj+" column: "+str(c)+", row: " +str(r))
                tableDocx = doc.tables[coordinatesDocX[i][4]]
                targetCell = tableDocx.cell(r, c)

                return targetCell

    #get name
    pupil = 0
    for name in names:

        #open .docx
        print(inputDocx)
        outputDocx = "output/"+name+"-saved-"+inputDocx
        savepath= outputDocx
        doc = docx.Document(inputDocx)

        #could make function getNote for the following 3 steps:
        #get note from allGrades
        for schoolSubject in allGrades:
            note = allGrades[schoolSubject]["Gesamtnote"][pupil]
            targetCell = getCoordinates(schoolSubject)
            if targetCell != None:
                #write to target
                targetCell.text = str(note)

        #get note from allGrades
        noteMitarbeit = allGrades["7a - Deu"]["Ø Mitarbeit"][pupil]
        noteKurztest = allGrades["7a - Deu"]["Ø Kurztest"][pupil]
        if noteMitarbeit != "x" and noteKurztest != "x":
            note = (float(noteMitarbeit.replace(",", ".")) + float(noteKurztest.replace(",", "."))) /2
        else:
            if noteMitarbeit != "x":
                note = noteMitarbeit
            else:
                note = "x"
        print("note: ", note)
        targetCell = getCoordinates("Deu_m")
        if targetCell != None:
            #write to target
            targetCell.text = str(note)

        #get note from allGrades
        note = allGrades["7a - Deu"]["Ø Klausur"][pupil]
        targetCell = getCoordinates("Deu_s")
        if targetCell != None:
            #write to target
            targetCell.text = str(note)


        pupil += 1

        #save to target .docX
        doc.save(savepath)
        print("Was saved in", savepath)


if __name__ == "__main__":

    #docxFile = "Z 420 - Notenzeugnis für Schülerinnen und Schüler mit dem Förderbedarf Lernen (01.21)-1.docx"
    #docxFile = "Zwischenbericht.docx"
    docxFile = "Zwischenbericht1.docx"
    #docxFile = "Z 600 - Arbeits- und Sozialverhalten.docx"
    xlsxFile = "/home/alex/schoolReport/Notenliste-18-10-2021_08-34-50.xlsx"
    FehlzeitenxlsxFile = "/home/alex/schoolReport/Unterrichtsstunden-18-10-2021_15-15-57.xlsx"


    writeSchoolReport(docxFile)

    names, allGrades = testDictionary.testDictionary(xlsxFile)
    pupil = 0
    for name in names:
        print(name)
        outputDocx = "output/"+name+"-saved-"+docxFile
        savepath= outputDocx

        testwritePupilDetailsToSchoolReport.testwritePupilDetailsToSchoolReport(inputDocx=savepath,
                                                                        docxFile=docxFile,
                                                                        nameOfPupil= name,
                                                                        save=savepath)
        writeParticipationToSchoolReport.writeParticipationToSchoolReport(inputDocx=savepath,
                                                                        docxFile=docxFile,
                                                                        nameOfPupil= name,
                                                                        save=savepath,
                                                                        xlsxFile=FehlzeitenxlsxFile,
                                                                        pupil=pupil)
        pupil += 1
